from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Risk Score"

    header_fill = PatternFill("solid", fgColor="DCE6F1")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    ws["A1"] = "ปัจจัย"
    ws["B1"] = "ผลตรวจ"
    ws["C1"] = "คะแนนความรุนแรง"
    ws["D1"] = "ค่าน้ำหนัก"
    ws["E1"] = "คะแนนถ่วงน้ำหนัก"
    ws["A14"] = "คำอธิบาย"
    ws["A15"] = "Workbook นี้ใช้เพื่ออธิบายตรรกะการคำนวณคำแนะนำและข้อความคาดการณ์ความเสี่ยงสุขภาพ"
    ws["A16"] = "1. แปลงผลตรวจของทั้ง 7 ปัจจัยเป็นคะแนนความรุนแรง"
    ws["A17"] = "2. นำคะแนนความรุนแรงคูณค่าน้ำหนักของแต่ละปัจจัย"
    ws["A18"] = "3. รวมคะแนนถ่วงน้ำหนักทั้งหมด และตรวจคะแนนสูงสุดของแต่ละปัจจัย"
    ws["A19"] = "4. นับจำนวนระบบสำคัญที่ผิดปกติร่วมกัน ได้แก่ ความดัน น้ำตาล ไต ตับ และไขมัน"
    ws["A20"] = "5. ใช้คะแนนรวม คะแนนสูงสุด และจำนวนระบบสำคัญ เพื่อสรุปคำแนะนำและข้อความคาดการณ์"

    factors = [
        ("ดัชนีมวลกาย (BMI)", 1.0, "อ้วน"),
        ("ความดันโลหิต", 1.3, "ความดันสูง"),
        ("ระดับน้ำตาลในเลือด", 1.4, "เบาหวาน"),
        ("ไขมันในเลือด", 1.1, "เฝ้าระวัง"),
        ("การทำงานของไต", 1.3, "ปกติ"),
        ("การทำงานของตับ", 1.2, "ปกติ"),
        ("กรดยูริกในเลือด", 0.8, "สูงกว่าปกติ"),
    ]

    for row_idx, (factor, weight, example_result) in enumerate(factors, start=2):
        ws[f"A{row_idx}"] = factor
        ws[f"B{row_idx}"] = example_result
        ws[f"D{row_idx}"] = weight
        ws[f"E{row_idx}"] = f"=C{row_idx}*D{row_idx}"

    ws["C2"] = '=IF(B2="ปกติ",0,IF(OR(B2="น้ำหนักเกิน",B2="ผอม"),1,IF(B2="อ้วน",2,IF(B2="อ้วนอันตราย",3,0))))'
    ws["C3"] = '=IF(B3="ปกติ",0,IF(B3="ความดันต่ำ",1,IF(B3="ความดันสูง",2,0)))'
    ws["C4"] = '=IF(B4="ปกติ",0,IF(B4="เสี่ยงเบาหวาน",2,IF(B4="เบาหวาน",3,IF(B4="ต่ำกว่าปกติ",1,0))))'
    ws["C5"] = '=IF(B5="ปกติ",0,IF(B5="เฝ้าระวัง",1,IF(B5="เสี่ยงต่อสุขภาพ",2,0)))'
    ws["C6"] = '=IF(B6="ปกติ",0,IF(B6="ผิดปกติเล็กน้อย",1,IF(B6="ผิดปกติปานกลาง",2,IF(B6="ผิดปกติรุนแรง",3,IF(B6="สูงกว่าปกติ",1,IF(B6="ต่ำกว่าปกติ",1,0))))))'
    ws["C7"] = '=IF(B7="ปกติ",0,IF(B7="ผิดปกติเล็กน้อย",1,IF(B7="ผิดปกติปานกลาง",2,IF(B7="ผิดปกติรุนแรง",3,0))))'
    ws["C8"] = '=IF(B8="ปกติ",0,IF(OR(B8="สูงกว่าปกติ",B8="ต่ำกว่าปกติ"),1,0))'

    ws["G1"] = "รายการสรุป"
    ws["H1"] = "ค่า"
    ws["G2"] = "คะแนนรวมถ่วงน้ำหนัก"
    ws["H2"] = "=SUM(E2:E8)"
    ws["G3"] = "คะแนนความรุนแรงสูงสุด"
    ws["H3"] = "=MAX(C2:C8)"
    ws["G4"] = "จำนวนระบบสำคัญที่ผิดปกติ"
    ws["H4"] = '=COUNTIF(C3:C7,">0")'
    ws["G5"] = "คำแนะนำ"
    ws["H5"] = '=IF(H3>=3,"ควรพบแพทย์อย่างเร่งด่วน",IF(AND(H4>=2,H2>=4.5),"ควรพบแพทย์อย่างเร่งด่วน",IF(H4>=2,"ควรพบแพทย์",IF(H2>=6.5,"ควรพบแพทย์อย่างเร่งด่วน",IF(H2>=3,"ควรพบแพทย์",IF(H2>0,"ควรปรึกษาแพทย์","ปกติ"))))))'
    ws["G6"] = "ข้อความคาดการณ์"
    ws["H6"] = '=IF(H3>=3,"หากไม่เข้ารับการดูแลทันที อาจเกิดภาวะแทรกซ้อนรุนแรง เช่น โรคหัวใจ หลอดเลือด หรือการเสื่อมของอวัยวะสำคัญ",IF(AND(H4>=2,H2>=4.5),"หากไม่เข้ารับการดูแลทันที อาจเกิดภาวะแทรกซ้อนรุนแรง เช่น โรคหัวใจ หลอดเลือด หรือการเสื่อมของอวัยวะสำคัญ",IF(H4>=2,"หากไม่ติดตามรักษาอย่างต่อเนื่อง มีโอกาสเพิ่มความเสี่ยงโรคเรื้อรังและภาวะแทรกซ้อนของระบบเมตาบอลิก",IF(H2>=6.5,"หากไม่เข้ารับการดูแลทันที อาจเกิดภาวะแทรกซ้อนรุนแรง เช่น โรคหัวใจ หลอดเลือด หรือการเสื่อมของอวัยวะสำคัญ",IF(H2>=3,"หากไม่ติดตามรักษาอย่างต่อเนื่อง มีโอกาสเพิ่มความเสี่ยงโรคเรื้อรังและภาวะแทรกซ้อนของระบบเมตาบอลิก",IF(H2>0,"หากไม่ปรับพฤติกรรมและติดตามผล อาจพัฒนาไปสู่ภาวะเสี่ยงโรคเรื้อรังในระยะถัดไป","ความเสี่ยงรวมอยู่ในระดับต่ำ หากดูแลสุขภาพต่อเนื่องตามปกติ"))))))'

    ws["G8"] = "เกณฑ์คะแนน"
    ws["G9"] = "ปกติ = 0"
    ws["G10"] = "เฝ้าระวัง/เล็กน้อย = 1"
    ws["G11"] = "ปานกลาง = 2"
    ws["G12"] = "รุนแรง = 3"

    for cell_range in ("A1:E1", "G1:H1"):
        for row in ws[cell_range]:
            for cell in row:
                cell.font = bold
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2, max_row=8, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            if cell.column in (3, 4, 5):
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2, max_row=6, min_col=7, max_col=8):
        for cell in row:
            cell.border = border
            if cell.column == 8 and cell.row != 6:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws["G6"].alignment = Alignment(vertical="top")
    ws["H6"].alignment = Alignment(wrap_text=True, vertical="top")

    for cell in ("D2", "D3", "D4", "D5", "D6", "D7", "D8", "E2", "E3", "E4", "E5", "E6", "E7", "E8", "H2"):
        ws[cell].number_format = "0.0"

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["G"].width = 28
    ws.column_dimensions["H"].width = 90
    ws.row_dimensions[6].height = 72
    ws.row_dimensions[15].height = 26
    ws.row_dimensions[16].height = 22
    ws.row_dimensions[17].height = 22
    ws.row_dimensions[18].height = 22
    ws.row_dimensions[19].height = 24
    ws.row_dimensions[20].height = 24

    ws["A14"].font = bold
    ws["A14"].fill = header_fill
    ws["A14"].border = border
    ws["A15"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A16"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A17"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A18"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A19"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A20"].alignment = Alignment(wrap_text=True, vertical="top")

    output_path = "health-risk-weighted-score.xlsx"
    wb.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
